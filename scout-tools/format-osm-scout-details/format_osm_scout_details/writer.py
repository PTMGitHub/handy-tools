from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from format_osm_scout_details.transform import OutputRow

TOTAL_COLUMNS = 7  # A..G

HEADERS = ["Name", "DOB", "Patrol", "Role", "Address", "Parents/Guardians", "Photo?"]

COLUMN_WIDTHS = {"A": 22, "B": 15, "C": 12, "D": 6, "E": 40, "F": 44, "G": 8}

WRAP_COLUMNS = {5, 6}  # Address, Parents/Guardians

HEADER_FILL = PatternFill("solid", fgColor="BFBFBF")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
NO_FILL = PatternFill(fill_type=None)
DIVIDER = Side(style="thin", color="000000")


def _parents_text(row: OutputRow) -> str:
    lines = []
    if row.p1_name:
        lines.append(f"{row.p1_name} {row.p1_phone}".strip())
    if row.p2_name:
        lines.append(f"{row.p2_name} {row.p2_phone}".strip())
    return "\n".join(lines)


def _style_header(ws) -> None:
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="top")
        cell.fill = HEADER_FILL


def _write_scout_block(ws, first_row: int, row: OutputRow, shaded: bool) -> None:
    identity_row, info_row = first_row, first_row + 1
    fill = BAND_FILL if shaded else NO_FILL

    ws.cell(identity_row, 1, row.name)
    ws.cell(identity_row, 2, row.dob)
    ws.cell(identity_row, 3, row.patrol_name)
    ws.cell(identity_row, 4, row.role)
    ws.cell(identity_row, 5, row.address)
    ws.cell(identity_row, 6, _parents_text(row))
    ws.cell(identity_row, 7, row.photo)

    ws.cell(info_row, 1, row.info)
    ws.merge_cells(start_row=info_row, start_column=1, end_row=info_row, end_column=TOTAL_COLUMNS)

    for col in range(1, TOTAL_COLUMNS + 1):
        identity_cell = ws.cell(identity_row, col)
        info_cell = ws.cell(info_row, col)
        identity_cell.alignment = Alignment(vertical="top", wrap_text=col in WRAP_COLUMNS)
        info_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        identity_cell.fill = fill
        info_cell.fill = fill

    # Bottom border under the Info row closes off each Scout's block.
    for col in range(1, TOTAL_COLUMNS + 1):
        ws.cell(info_row, col).border = Border(bottom=DIVIDER)


def write_output(rows: List[OutputRow], path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scout Details"

    _style_header(ws)

    next_row = 2
    for index, row in enumerate(rows):
        _write_scout_block(ws, next_row, row, shaded=(index % 2 == 0))
        next_row += 2

    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(path)
