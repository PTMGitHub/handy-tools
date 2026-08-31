from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional

import openpyxl

WORKSHEET_NAME = "Scouts"

# OSM normally exports Date of birth as a real Excel date, but some exports
# (or cells that got re-typed by hand) store it as text instead -- these are
# the formats we've seen in practice.
DOB_TEXT_FORMATS = ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y")

GROUP_MEMBER = "Member"
GROUP_P1 = "Primary Contact 1"
GROUP_P2 = "Primary Contact 2"
GROUP_ESSENTIAL = "Essential Information"
GROUP_ADDITIONAL = "Additional Information"
GROUP_SURGERY = "Doctor's Surgery"
GROUP_CONSENTS = "Consents"


@dataclass
class ScoutRecord:
    first_name: str
    last_name: str
    known_as: str
    dob: Optional[date]
    patrol_name: str
    patrol_role: str

    member_address_parts: List[str] = field(default_factory=list)

    p1_first_name: str = ""
    p1_last_name: str = ""
    p1_relationship: str = ""
    p1_address_parts: List[str] = field(default_factory=list)
    p1_phone_1: str = ""
    p1_phone_2: str = ""

    p2_first_name: str = ""
    p2_last_name: str = ""
    p2_relationship: str = ""
    p2_phone_1: str = ""
    p2_phone_2: str = ""

    medical_details: str = ""
    allergies: str = ""
    dietary_requirements: str = ""
    other_useful_information: str = ""
    surgery_medical_information: str = ""
    surgery_dietary_needs: str = ""

    photographs: str = ""


def _build_column_lookup(ws):
    """Map (group, field name) -> 1-based column index from the export's
    two header rows. The group (row 1) is forward-filled across the blank
    cells that follow it, matching how OSM merges group headers visually.
    """
    row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    lookup = {}
    group = None
    for idx, (group_cell, header) in enumerate(zip(row1, row2), start=1):
        if group_cell:
            group = group_cell
        lookup.setdefault((group, header), idx)
    return lookup


def _cell(ws, row, lookup, group, header):
    col = lookup.get((group, header))
    if col is None:
        return ""
    value = ws.cell(row=row, column=col).value
    if isinstance(value, str):
        return value.strip()
    return value if value is not None else ""


def _read_dob(ws, row, lookup, group, header) -> Optional[date]:
    """Parse Date of birth defensively: OSM normally exports it as a real
    Excel date, but a blank cell (seen for some adult/helper records) or a
    cell that's actually text (e.g. re-typed by hand) both show up here, and
    neither should crash the whole run.
    """
    value = _cell(ws, row, lookup, group, header)
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in DOB_TEXT_FORMATS:
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def read_scouts(path: Path) -> List[ScoutRecord]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[WORKSHEET_NAME]
    lookup = _build_column_lookup(ws)

    records = []
    row = 3
    while True:
        first_name = _cell(ws, row, lookup, None, "First name")
        if not first_name:
            break

        records.append(
            ScoutRecord(
                first_name=first_name,
                last_name=_cell(ws, row, lookup, None, "Last name"),
                known_as=_cell(ws, row, lookup, GROUP_ADDITIONAL, "Known As"),
                dob=_read_dob(ws, row, lookup, None, "Date of birth"),
                patrol_name=_cell(ws, row, lookup, None, "Patrol Name"),
                patrol_role=_cell(ws, row, lookup, None, "Patrol Role"),
                member_address_parts=[
                    _cell(ws, row, lookup, GROUP_MEMBER, "Address 1"),
                    _cell(ws, row, lookup, GROUP_MEMBER, "Address 2"),
                    _cell(ws, row, lookup, GROUP_MEMBER, "Address 3"),
                    _cell(ws, row, lookup, GROUP_MEMBER, "Address 4"),
                    _cell(ws, row, lookup, GROUP_MEMBER, "Postcode"),
                ],
                p1_first_name=_cell(ws, row, lookup, GROUP_P1, "First Name"),
                p1_last_name=_cell(ws, row, lookup, GROUP_P1, "Last Name"),
                p1_relationship=_cell(ws, row, lookup, GROUP_P1, "Relationship"),
                p1_address_parts=[
                    _cell(ws, row, lookup, GROUP_P1, "Address 1"),
                    _cell(ws, row, lookup, GROUP_P1, "Address 2"),
                    _cell(ws, row, lookup, GROUP_P1, "Address 3"),
                    _cell(ws, row, lookup, GROUP_P1, "Address 4"),
                    _cell(ws, row, lookup, GROUP_P1, "Postcode"),
                ],
                p1_phone_1=_cell(ws, row, lookup, GROUP_P1, "Phone 1"),
                p1_phone_2=_cell(ws, row, lookup, GROUP_P1, "Phone 2"),
                p2_first_name=_cell(ws, row, lookup, GROUP_P2, "First Name"),
                p2_last_name=_cell(ws, row, lookup, GROUP_P2, "Last Name"),
                p2_relationship=_cell(ws, row, lookup, GROUP_P2, "Relationship"),
                p2_phone_1=_cell(ws, row, lookup, GROUP_P2, "Phone 1"),
                p2_phone_2=_cell(ws, row, lookup, GROUP_P2, "Phone 2"),
                medical_details=_cell(ws, row, lookup, GROUP_ESSENTIAL, "Medical details"),
                allergies=_cell(ws, row, lookup, GROUP_ESSENTIAL, "Allergies"),
                dietary_requirements=_cell(ws, row, lookup, GROUP_ESSENTIAL, "Dietary requirements"),
                other_useful_information=_cell(ws, row, lookup, GROUP_ESSENTIAL, "Other useful information"),
                surgery_medical_information=_cell(ws, row, lookup, GROUP_SURGERY, "Medical Information"),
                surgery_dietary_needs=_cell(ws, row, lookup, GROUP_SURGERY, "Dietary Needs"),
                photographs=_cell(ws, row, lookup, GROUP_CONSENTS, "Photographs"),
            )
        )
        row += 1

    return records
