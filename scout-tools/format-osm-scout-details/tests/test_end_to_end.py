from pathlib import Path

import openpyxl

from format_osm_scout_details.reader import read_scouts
from format_osm_scout_details.transform import sort_records, transform
from format_osm_scout_details.writer import write_output

REFERENCE_DIR = Path(__file__).resolve().parent.parent / "reference_files"


def _read_scout_blocks(ws, num_scouts):
    """Each Scout occupies two physical rows starting at row 2: an identity
    row (Name, DOB, Patrol, Role, Address, Parents/Guardians, Photo?) and an
    Info row merged across all columns."""
    blocks = []
    for i in range(num_scouts):
        identity_row = 2 + i * 2
        info_row = identity_row + 1
        values = [c.value for c in ws[identity_row]]
        blocks.append(
            dict(
                name=values[0], dob=values[1], patrol_name=values[2], role=values[3],
                address=values[4], parents=values[5], photo=values[6],
                info=ws.cell(info_row, 1).value,
            )
        )
    return blocks


def test_pipeline_matches_example_output(tmp_path):
    records = read_scouts(REFERENCE_DIR / "example_input.xlsx")
    records = sort_records(records)
    rows = [transform(record) for record in records]

    out_path = tmp_path / "actual_output.xlsx"
    write_output(rows, out_path)

    actual_ws = openpyxl.load_workbook(out_path, data_only=True)["Scout Details"]
    expected_ws = openpyxl.load_workbook(
        REFERENCE_DIR / "example_output.xlsx", data_only=True
    )["Scout Details"]

    actual_blocks = _read_scout_blocks(actual_ws, len(rows))
    expected_blocks = _read_scout_blocks(expected_ws, len(rows))

    assert actual_blocks == expected_blocks
